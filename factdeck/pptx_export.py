"""Draw measured HTML slots as native PowerPoint shapes. Text stays text."""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from factdeck.slots import measured_text

SLIDE_W = 12_192_000
SLIDE_H = 6_858_000
PX_TO_EMU = SLIDE_W / 1920
FONT_NAME = "WenQuanYi Micro Hei"
PAPER = "FFFFFF"

_COLOR = re.compile(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,\s*([0-9.]+))?\s*\)")


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    return prs


def _px(value) -> float:
    if value is None:
        return 0.0
    return float(str(value).replace("px", "").strip() or 0)


def _color(value: str | None) -> RGBColor | None:
    if not value:
        return None
    match = _COLOR.match(value.strip())
    if not match:
        return None
    if match.group(4) is not None and float(match.group(4)) < 0.05:
        return None
    return RGBColor(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def _emu(px: float) -> Emu:
    return Emu(int(round(px * PX_TO_EMU)))


def _set_run_font(run, size_px: float, color: RGBColor | None, bold: bool) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(max(8, size_px * 0.5))
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        element = r_pr.find(qn(tag))
        if element is None:
            element = r_pr.makeelement(qn(tag), {})
            r_pr.append(element)
        element.set("typeface", FONT_NAME)


def _align(value: str):
    if value == "center":
        return PP_ALIGN.CENTER
    if value == "right":
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def add_measured_slide(prs: Presentation, slots: list[dict], paper: str = PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(paper)
    ordered = [slot for slot in slots if slot.get("kind") in {"shape", "line"}]
    ordered += [slot for slot in slots if slot.get("kind") == "text"]
    for slot in ordered:
        _draw(slide, slot)
    return slide


def _draw(slide, slot: dict) -> None:
    kind = slot.get("kind")
    width = float(slot.get("w") or 0)
    height = float(slot.get("h") or 0)
    if width < 0.5 or height < 0.5:
        return
    left = _emu(float(slot.get("x") or 0))
    top = _emu(float(slot.get("y") or 0))
    box_w = _emu(max(width, 1))
    box_h = _emu(max(height, 1))
    text = measured_text(slot)
    if kind == "line":
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, box_w, box_h)
        shape.shadow.inherit = False
        color = _color(slot.get("fill")) or RGBColor(0x5C, 0x6B, 0x7A)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return
    if kind == "shape":
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, box_w, box_h)
        shape.shadow.inherit = False
        fill = _color(slot.get("fill"))
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        _apply_border(shape, slot.get("border") or {})
        if text:
            _write_text(shape, slot, text)
        else:
            shape.text_frame.clear()
        return
    if not text:
        return
    fill = _color(slot.get("fill"))
    if fill is not None:
        back = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, box_w, box_h)
        back.shadow.inherit = False
        back.fill.solid()
        back.fill.fore_color.rgb = fill
        back.line.fill.background()
    box = slide.shapes.add_textbox(left, top, box_w, _emu(height + 2))
    box.shadow.inherit = False
    _write_text(box, slot, text)
    _apply_border(box, slot.get("border") or {})


def _apply_border(shape, border: dict) -> None:
    sides = [border.get(name) or {} for name in ("top", "right", "bottom", "left")]
    solid = [side for side in sides if side.get("style") == "solid" and _px(side.get("w")) >= 1]
    dashed = [side for side in sides if side.get("style") == "dashed" and _px(side.get("w")) >= 1]
    if len(solid) >= 4:
        shape.line.color.rgb = _color(solid[0].get("color")) or RGBColor(0x0E, 0x2C, 0x4A)
        shape.line.width = Pt(max(0.75, _px(solid[0].get("w")) * 0.5))
        return
    if len(dashed) >= 4:
        shape.line.color.rgb = _color(dashed[0].get("color")) or RGBColor(0x8A, 0x81, 0x75)
        shape.line.width = Pt(max(0.75, _px(dashed[0].get("w")) * 0.5))
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return
    shape.line.fill.background()
    left = border.get("left") or {}
    if left.get("style") == "solid" and _px(left.get("w")) >= 4:
        # Accent bar is already inside the measured shape box via the CSS border.
        # A full outline would double it; the HTML border-left is not a separate slot.
        # Draw it by insetting a filled rectangle on the left edge.
        color = _color(left.get("color")) or RGBColor(0x1F, 0x4E, 0x79)
        bar_w = _emu(_px(left.get("w")))
        bar = shape._parent.add_shape(MSO_SHAPE.RECTANGLE, shape.left, shape.top, bar_w, shape.height)
        bar.shadow.inherit = False
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()


def _write_text(shape, slot: dict, text: str) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    try:
        frame._txBody.bodyPr.set("anchor", "t")
    except Exception:
        frame.paragraphs[0].alignment = _align(slot.get("align") or "left")
    lines = text.splitlines() or [text]
    color = _color(slot.get("color"))
    size = _px(slot.get("size") or 18)
    bold = int(float(slot.get("weight") or 400)) >= 600
    align = _align(slot.get("align") or "left")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        _set_run_font(run, size, color, bold)
    anchor = MSO_ANCHOR.TOP
    frame.paragraphs[0].alignment = align
    _ = anchor
